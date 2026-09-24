from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from meta_ads_intelligence.pipeline import run


class PipelineTest(unittest.TestCase):
    def test_deduplicates_and_tracks_history_without_marking_missing_ads(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_dir = root / "imports" / "2026-09-24"
            input_dir.mkdir(parents=True)
            pd.DataFrame([
                {"ad_id": "100", "page_id": "0012", "ad_text": "Alpha", "ad_link": "https://example.com/a"},
                {"ad_id": "100", "page_id": "0012", "ad_text": "Alpha duplicate", "ad_link": "https://example.com/a"},
                {"ad_id": "200", "page_id": "0012", "ad_text": "Beta"},
            ]).to_csv(input_dir / "one.csv", index=False)
            pd.DataFrame([{"ad_id": "300", "page_id": "0012", "ad_text": "Gamma"}]).to_csv(input_dir / "two.csv", index=False)

            report = run(input_dir, root / "reports", root / "archive", scan_date="2026-09-24")
            workbook = load_workbook(report)
            sheet = workbook["الإعلانات"]
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet["C2"].value, "100")
            self.assertEqual(sheet["C2"].data_type, "s")
            self.assertEqual(sheet["N2"].value, "جديد في سجلنا")
            self.assertTrue(sheet["K2"].hyperlink)
            latest = json.loads((root / "docs" / "data" / "latest.json").read_text(encoding="utf-8"))
            self.assertEqual(latest["scan_date"], "2026-09-24")
            self.assertEqual(latest["total_ads"], 3)
            self.assertEqual(latest["total_new_ads"], 3)
            self.assertEqual({"ad_id", "page_id", "page_name", "comparison_result"} - set(latest["ads"][0]), set())
            self.assertEqual(latest["competitors"][0]["total_ads"], 3)

            second_input = root / "imports" / "2026-10-01"
            second_input.mkdir(parents=True)
            pd.DataFrame([{"ad_id": "100", "page_id": "0012", "ad_text": "Alpha"}]).to_csv(second_input / "scan.csv", index=False)
            second_report = run(second_input, root / "reports", root / "archive", scan_date="2026-10-01")
            second_sheet = load_workbook(second_report)["الإعلانات"]
            self.assertEqual(second_sheet["N2"].value, "موجود سابقاً")
            self.assertEqual(second_sheet["Q2"].value, 2)
            history = json.loads((root / "archive" / "history.json").read_text(encoding="utf-8"))
            self.assertEqual(set(history), {"100", "200", "300"})
            self.assertNotIn("stopped", history["200"])
            latest_after_second_scan = json.loads((root / "docs" / "data" / "latest.json").read_text(encoding="utf-8"))
            self.assertEqual(latest_after_second_scan["total_ads"], 1)
            self.assertEqual(latest_after_second_scan["total_new_ads"], 0)


if __name__ == "__main__":
    unittest.main()
