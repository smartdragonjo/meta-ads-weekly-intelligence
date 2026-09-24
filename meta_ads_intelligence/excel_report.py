from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


AD_HEADERS = [
    ("competitor_name", "اسم المنافس"),
    ("page_id", "معرف الصفحة"),
    ("ad_id", "رقم الإعلان"),
    ("ad_text", "نص الإعلان"),
    ("ad_title", "عنوان الإعلان"),
    ("start_date", "تاريخ بدء الإعلان"),
    ("status", "حالة الإعلان"),
    ("ad_type", "نوع الإعلان"),
    ("cta", "CTA"),
    ("platforms", "المنصات"),
    ("ad_link", "رابط الإعلان"),
    ("video_link", "رابط الفيديو"),
    ("image_link", "رابط الصورة"),
    ("comparison", "نتيجة المقارنة"),
    ("first_seen", "أول ظهور"),
    ("last_seen", "آخر ظهور"),
    ("times_seen", "عدد مرات الظهور"),
]

SUMMARY_HEADERS = [
    "تاريخ الفحص", "اسم المنافس", "معرف الصفحة", "عدد الإعلانات الفريدة",
    "عدد الجديدة", "عدد الموجودة سابقاً", "ملاحظة",
]

LINK_FIELDS = {"ad_link", "video_link", "image_link"}


def _set_link(cell, value: object) -> None:
    text = str(value or "").strip()
    if text.startswith(("http://", "https://")):
        cell.hyperlink = text
        cell.style = "Hyperlink"


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.rightToLeft = True
    sheet.row_dimensions[1].height = 28
    header_fill = PatternFill("solid", fgColor="16324F")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_report(path: Path, ads: pd.DataFrame, scan_date: str, competitor_name: str = "", page_id: str = "") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    export = ads.copy()
    if competitor_name:
        export["competitor_name"] = competitor_name
    if page_id:
        export["page_id"] = page_id

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        export[[field for field, _ in AD_HEADERS]].rename(
            columns={field: header for field, header in AD_HEADERS}
        ).to_excel(writer, sheet_name="الإعلانات", index=False)

        new_count = int((export["comparison"] == "جديد في سجلنا").sum())
        existing_count = int((export["comparison"] == "موجود سابقاً").sum())
        summary = pd.DataFrame([[
            scan_date,
            competitor_name,
            page_id,
            len(export),
            new_count,
            existing_count,
            "البيانات عينة وليست دليلاً على توقف إعلان غائب.",
        ]], columns=SUMMARY_HEADERS)
        summary.to_excel(writer, sheet_name="ملخص التشغيل", index=False)

    workbook = load_workbook(path)
    ads_sheet = workbook["الإعلانات"]
    summary_sheet = workbook["ملخص التشغيل"]
    _style_sheet(ads_sheet, [20, 18, 20, 48, 32, 18, 16, 16, 16, 22, 34, 34, 34, 22, 16, 16, 18])
    _style_sheet(summary_sheet, [18, 24, 20, 22, 16, 24, 58])

    field_indexes = {field: index + 1 for index, (field, _) in enumerate(AD_HEADERS)}
    for row in ads_sheet.iter_rows(min_row=2):
        for field in LINK_FIELDS:
            _set_link(row[field_indexes[field] - 1], row[field_indexes[field] - 1].value)

    workbook.save(path)
